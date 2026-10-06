"""Универсальный импорт данных завода: Word (.docx), Excel (.xlsx), CSV.

Реальные выгрузки будут отличаться от тестового файла ТЗ: другие названия колонок, порядок, лишние столбцы,
даты как в Excel. Поэтому таблицы узнаём ПО СМЫСЛУ заголовков (синонимы RU/EN), а не по позиции.
Если чего-то не хватает — понятная ошибка, а не тихо искажённые числа.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path

from .ingest_docx import ASSET_IDS, ROUTE, IngestError, _parse_targets
from .schemas import DowntimeRecord, Facts, LineRecord, ModelPlan, Provenance, QualityRecord, Targets

# синонимы заголовков → поле
SYN = {
    "date": ["дата", "date", "день", "дата смены", "shift date"],
    "line": ["линия", "line", "производственная линия"],
    "area": ["участок", "цех", "area", "section", "workshop", "station"],
    "plan": ["план", "plan", "план шт", "planned"],
    "fact": ["факт", "fact", "actual", "выпуск факт", "факт шт"],
    "runtime": ["время работы ч", "время работы", "runtime", "часы работы", "работа ч", "run time h"],
    "util": ["загрузка", "utilization", "load", "загрузка %"],
    "equipment": ["оборудование", "equipment", "станок", "asset", "машина"],
    "reason": ["причина", "reason", "cause", "причина простоя"],
    "minutes": ["длительность мин", "длительность", "minutes", "простой мин", "duration", "duration min", "мин"],
    "model": ["модель", "model", "марка"],
    "monthly": ["план на месяц", "monthly plan", "план мес", "месячный план"],
    "produced": ["выпущено", "produced", "выпуск", "количество", "output"],
    "defects": ["брак", "defects", "дефекты", "брак шт", "rejects"],
    "defect_pct": ["% брака", "брака %", "defect %", "процент брака", "defect rate"],
}
AREA_WORDS = {"welding": ["свар", "weld", "body"], "paint": ["окрас", "краск", "paint"], "assembly": ["сбор", "assembl", "конвейер"]}
DEFAULT_TARGETS = Targets(shifts_per_day=2, shift_hours=8, oee_min_pct=85, defect_max_pct=2,
                          critical_downtime_max_min_per_day=60, monthly_output_min=5500)


def _norm(h) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[,.;:()\"']", " ", str(h or "").lower())).strip()


def _field(header) -> str | None:
    h = _norm(header)
    if not h:
        return None
    best = None
    for f, words in SYN.items():
        for w in words:
            w = _norm(w)
            if h == w:
                return f
            if w in h and (best is None or len(w) > best[1]):
                best = (f, len(w))
    return best[0] if best else None


def _num(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace(",", ".").replace("%", "")
    if not s:
        raise ValueError("пусто")
    return float(s)


def _date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y", "%Y.%m.%d"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    raise ValueError(f"не похоже на дату: «{s}»")


def _area(text) -> str:
    t = str(text).lower()
    for a, words in AREA_WORDS.items():
        if any(w in t for w in words):
            return a
    raise ValueError(f"не узнан участок «{text}» (ожидается сварка / окраска / сборка)")


def _kind(fields: set[str]) -> str | None:
    if {"produced", "defects"} <= fields:
        return "quality"
    if "minutes" in fields and ({"equipment", "reason"} & fields):
        return "downtime"
    if "model" in fields and ({"monthly", "plan"} & fields):
        return "model_plan"
    if {"plan", "fact"} <= fields:
        return "lines"
    return None


# ── чтение файлов в «сырые таблицы» ──────────────────────────────────────────
def _tables_docx(data: bytes) -> tuple[list[list[list]], str]:
    import docx
    d = docx.Document(io.BytesIO(data))
    tables = [[[c.text.strip() for c in r.cells] for r in t.rows] for t in d.tables]
    text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
    return tables, text


def _tables_xlsx(data: bytes) -> tuple[list[list[list]], str]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    tables, texts = [], []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        texts += [str(c) for r in rows for c in r if isinstance(c, str)]
        # заголовок — первая строка (из первых 10), где узнано ≥2 поля
        for i, r in enumerate(rows[:10]):
            if sum(1 for c in r if _field(c)) >= 2:
                tables.append(rows[i:])
                break
    return tables, "\n".join(texts)


def _tables_csv(data: bytes) -> tuple[list[list[list]], str]:
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    return [rows], ""


def read_tables(name: str, data: bytes) -> tuple[list[list[list]], str]:
    ext = Path(name).suffix.lower()
    if ext == ".docx":
        return _tables_docx(data)
    if ext in (".xlsx", ".xlsm"):
        return _tables_xlsx(data)
    if ext in (".csv", ".txt"):
        return _tables_csv(data)
    raise IngestError(f"Формат {ext} не поддерживается: загрузите .docx, .xlsx или .csv")


# ── сырые таблицы → факты ────────────────────────────────────────────────────
def facts_from_tables(name: str, tables: list[list[list]], text: str) -> tuple[Facts, list[str]]:
    lines, downtime, model_plan, quality = [], [], [], []
    warnings: list[str] = []
    for ti, rows in enumerate(tables, start=1):
        if len(rows) < 2:
            continue
        cols = {}
        for j, h in enumerate(rows[0]):
            f = _field(h)
            if f and f not in cols:
                cols[f] = j
        kind = _kind(set(cols))
        if kind is None:
            warnings.append(f"Таблица {ti}: не распознана (заголовки: {', '.join(str(h) for h in rows[0] if h)}) — пропущена")
            continue
        for ri, r in enumerate(rows[1:], start=1):
            g = lambda f: r[cols[f]] if f in cols and cols[f] < len(r) else None  # noqa: E731
            prov = Provenance(file=name, table=ti, row=ri)
            try:
                if kind == "lines":
                    line = str(g("line") or g("area") or "")
                    plan, fact = int(_num(g("plan"))), int(_num(g("fact")))
                    rt = _num(g("runtime")) if g("runtime") not in (None, "") else None
                    ut = _num(g("util")) if g("util") not in (None, "") else (round(rt / 8 * 100) if rt is not None else None)
                    if rt is None:
                        warnings.append(f"Таблица {ti}, стр. {ri}: нет времени работы — принято 8 ч (допущение)")
                        rt = 8.0
                    lines.append(LineRecord(date=_date(g("date")), line=line, area=_area(line), plan=plan, fact=fact,
                                            runtime_h=rt, utilization_pct=ut if ut is not None else 100, provenance=prov))
                elif kind == "downtime":
                    eq = str(g("equipment") or "—")
                    reason = str(g("reason") or "—")
                    downtime.append(DowntimeRecord(date=_date(g("date")), area=_area(g("area") or eq), asset_id=ASSET_IDS.get(eq, eq.lower()),
                                                   equipment=eq, reason=reason, minutes=int(_num(g("minutes"))),
                                                   planned=any(w in reason.lower() for w in ("планов", "то ", "planned")) or reason.lower().strip() == "то",
                                                   provenance=prov))
                elif kind == "model_plan":
                    model_plan.append(ModelPlan(model=str(g("model")), monthly_plan=int(_num(g("monthly") if "monthly" in cols else g("plan"))), provenance=prov))
                else:
                    produced, defects = int(_num(g("produced"))), int(_num(g("defects")))
                    pct = _num(g("defect_pct")) if g("defect_pct") not in (None, "") else round(100 * defects / max(1, produced), 1)
                    quality.append(QualityRecord(date=_date(g("date")), area=_area(g("area") or g("line")), produced=produced,
                                                 defects=defects, defect_pct_reported=pct, provenance=prov))
            except (ValueError, TypeError) as e:
                warnings.append(f"Таблица {ti}, стр. {ri}: {e} — строка пропущена")
    missing = [n for n, v in (("работа линий (план/факт)", lines), ("качество (выпущено/брак)", quality)) if not v]
    if missing:
        raise IngestError("Не найдены обязательные таблицы: " + "; ".join(missing))
    try:
        targets = _parse_targets(text.splitlines())
    except IngestError:
        targets = DEFAULT_TARGETS
        warnings.append("Нормы (OEE, брак, лимит простоя, план) не найдены — взяты из ТЗ (допущение)")
    if not model_plan:
        warnings.append("Нет плана по моделям — сверка плана с целью недоступна")
    facts = Facts(source_file=name, route=ROUTE, lines=lines, downtime=downtime, model_plan=model_plan,
                  quality=quality, targets=targets)
    return facts, warnings


def load_any(name: str, data: bytes) -> tuple[Facts, list[str]]:
    tables, text = read_tables(name, data)
    return facts_from_tables(name, tables, text)
