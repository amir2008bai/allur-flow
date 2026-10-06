"""Замена тестовых данных реальными: другой формат (Excel), другие заголовки, другие даты и оборудование."""
import io
import random
from datetime import date, timedelta

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


def make_xlsx(paint_defect=0.03, days=10) -> bytes:
    rnd = random.Random(1)
    wb = Workbook()
    ws = wb.active; ws.title = "Lines"
    ws.append(["Комментарий: выгрузка MES"])               # мусорная строка над заголовком
    ws.append(["Date", "Line", "Plan, pcs", "Actual", "Run time h", "Utilization %", "Мастер"])
    d0 = date(2026, 9, 1)
    for i in range(days):
        for line in ("Сварка-1", "Окраска-1", "Сборка-1"):
            rt = round(rnd.uniform(7.0, 8.0), 1)
            ws.append([d0 + timedelta(days=i), line, 130, int(130 * rt / 8 * rnd.uniform(0.93, 1.0)), rt, round(rt / 8 * 100), "Иванов"])
    ws2 = wb.create_sheet("Downtime")
    ws2.append(["Дата", "Цех", "Equipment", "Cause", "Duration min"])
    ws2.append([d0, "Окраска", "Booth-1", "Засор форсунки", 35])
    ws2.append([d0 + timedelta(days=3), "Сборка", "Line-2 conveyor", "Сбой привода", 48])
    ws3 = wb.create_sheet("Quality")
    ws3.append(["Дата", "Участок", "Выпущено", "Дефекты"])  # без % брака — посчитается
    for i in range(days):
        for a, p in (("Сварка", 0.015), ("Окраска", paint_defect), ("Сборка", 0.01)):
            n = 125; ws3.append([d0 + timedelta(days=i), a, n, sum(rnd.random() < p for _ in range(n))])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def test_preview_real_excel():
    c = TestClient(app)
    r = c.post("/api/data/preview", files={"file": ("mes_export.xlsx", make_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["rows"]["lines"] == 30 and p["rows"]["quality"] == 30 and p["rows"]["downtime"] == 2
    assert p["date_from"] == "2026-09-01" and p["areas"] == ["assembly", "paint", "welding"]


def test_activate_recalibrates_and_reset_restores():
    c = TestClient(app)
    test_paint = c.get("/api/learn/calib").json()["current"]["mean"]
    st = c.post("/api/data/activate", files={"file": ("mes_export.xlsx", make_xlsx(paint_defect=0.08), "application/octet-stream")}).json()
    try:
        assert st["is_test"] is False and st["file"].endswith("mes_export.xlsx")
        cal = st["calibration"]
        assert cal["defect_pct"]["paint"] > 5                       # брак окраски из новых данных
        assert abs(cal["model_per_shift"] - cal["target_per_shift"]) < 6  # темп двойника подогнан под факт
        assert c.get("/api/findings").status_code == 200              # выводы строятся на новых данных
        assert c.get("/api/sim/run", params={"scenario": "base", "seed": 3}).status_code == 200
        assert c.get("/api/learn/calib").json()["current"]["mean"] > test_paint
    finally:
        back = c.post("/api/data/reset").json()
    assert back["is_test"] is True and back["file"] == "test_data.docx"


def test_bad_file_gives_clear_error():
    c = TestClient(app)
    r = c.post("/api/data/preview", files={"file": ("x.csv", "a;b\n1;2\n".encode(), "text/csv")})
    assert r.status_code == 422 and "Не найдены обязательные таблицы" in r.json()["detail"]
